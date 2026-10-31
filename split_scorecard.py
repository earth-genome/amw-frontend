#!/usr/bin/env -S uv run --script
# /// script
# requires-python = ">=3.10"
# dependencies = [
#     "openpyxl>=3.1",
# ]
# ///
"""
Split the mining scorecard workbook into the three policy-scoreboard CSVs
(English, Spanish, Portuguese).

Usage:
    uv run split_scorecard.py "./tmp/Mining_scorecards (do not edit this one).xlsx"
    uv run split_scorecard.py "./tmp/Mining_scorecards (do not edit this one).xlsx" -o ./out --prefix policy-scoreboard-data

Output (one per language; <out> defaults to
public/amazon-mining-policy-scoreboard/data/ next to this script, i.e. in the
repo root, regardless of the current working directory):
    <out>/<prefix>_en.csv
    <out>/<prefix>_es.csv
    <out>/<prefix>_pt.csv

Notes
-----
* Scores (the "XX_Evaluation" values) are always taken from the "Database"
  sheet (SCORES_SHEET below), which is the source of truth. The language
  sheets (SHEETS below) supply only the texts: Question, Justification and
  Sources (plus Dimension / Categories / Category ID).
* Country columns are located by their "XX_Evaluation" header, not by
  position, so reordering the countries in the workbook does not break the
  export. The CSV column order is always the one in COUNTRIES below.
* The Justification and Sources of each country are the two columns right
  after its "XX_Evaluation". Their headers are free text in each language
  ("Bolivia justification", "Justificación de Bolivia", "Fontes Bolívia"...),
  so they are identified by keyword (see BLOCK_KEYWORDS), in either order.
  If they cannot be identified, the script stops rather than guessing.
* Columns for country codes that are not listed in COUNTRIES are reported as
  a warning (they are not exported).
* Ambiguous headers (two columns matching the same field) stop the script.
* Rows are matched and sorted by question number.
* "Dimension" and "Categories" are always written in English (taken from the
  English sheet), as in the published CSVs.
* Numeric scores are written as 1 / 0.5 / 0. Any other number is written
  as-is and reported as a warning. Any non-numeric (textual) evaluation is
  written as-is (trimmed), and reported on the console. These checks are
  done on the "Database" sheet, since that is where the exported scores
  come from.
* The scores in each language sheet (including English) are cross-checked
  against the "Database" sheet; any mismatch is printed as a warning. The
  CSVs are still written, always with the "Database" scores.
* Questions present in the "Database" sheet but missing from a language
  sheet (or the other way round) are reported as a warning. A question that
  is in a language sheet but not in "Database" is exported with blank scores.
* With --strict, any warning (while reading or cross-checking) makes the
  script exit with code 1 after the CSVs are written.
"""

from __future__ import annotations

import argparse
import csv
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

# --------------------------------------------------------------------------- #
# Configuration
# --------------------------------------------------------------------------- #

# Language code -> worksheet name in the workbook.
SHEETS: dict[str, str] = {
    "en": "Translated english",
    "es": "Translated to spanish",
    "pt": "Translated to portuguese",
}

# The sheet that supplies the scores ("XX_Evaluation") for every CSV. It is
# the source of truth; the scores in the SHEETS above are only cross-checked
# against it.
SCORES_SHEET = "Database"

# Default output directory. The script lives in the repo root, so this is
# resolved relative to the script's own location, not the current directory.
REPO_ROOT = Path(__file__).resolve().parent
DEFAULT_OUTDIR = REPO_ROOT / "public" / "amazon-mining-policy-scoreboard" / "data"

# The sheet that supplies the (English) Dimension / Categories for every CSV.
REFERENCE_LANG = "en"

# Output order of countries: (code used in the "XX_Evaluation" header, CSV name)
COUNTRIES: list[tuple[str, str]] = [
    ("BO", "Bolivia"),
    ("BR", "Brazil"),
    ("CO", "Colombia"),
    ("PE", "Peru"),
    ("EC", "Ecuador"),
    ("GU", "Guyana"),
    ("SU", "Suriname"),
    ("VE", "Venezuela"),
]

# Header aliases (compared case-insensitively, whitespace-trimmed) for the
# fixed columns in each language sheet.
FIXED_COLUMNS: dict[str, set[str]] = {
    "Dimension": {"dimension", "dimensión", "dimensão"},
    "Categories": {"categories", "categorías", "categorias"},
    "Category ID": {"category id", "categoria id", "categoría id", "id de categoría"},
    "Question Number": {
        "question number",
        "número de la pregunta",
        "numero de la pregunta",
        "n. da pergunta",
        "número da pergunta",
    },
    "Question": {"question", "pregunta", "pergunta"},
}

# Accepted suffixes for the "XX_Evaluation" header (compared
# case-insensitively, whitespace-trimmed).
EVALUATION_SUFFIXES: set[str] = {"evaluation", "evaluación", "evaluacion",
                                 "avaliação", "avaliacao"}

# Keywords that identify the two columns after each "XX_Evaluation"
# (matched as substrings, case- and accent-insensitive). Both columns must be
# identified, exactly once each, or the script stops.
BLOCK_KEYWORDS: dict[str, tuple[str, ...]] = {
    "Justification": ("justif",),
    "Sources": ("source", "fuente", "fonte"),
}

# Valid numeric scores, as formatted by fmt_number(). Any other number is
# exported as-is but reported as a warning.
ALLOWED_SCORES: set[str] = {"0", "0.5", "1"}

# "BO_Evaluation", "BO Evaluation", "GUY_Evaluation", ...
EVAL_HEADER_RE = re.compile(r"^\s*([A-Za-z]{2,3})[ _](.+?)\s*$")

# Excel/OOXML escapes such as "_x000D_" or "_x0002_" that openpyxl leaves in
# read-only mode.
OOXML_ESCAPE_RE = re.compile(r"_x([0-9A-Fa-f]{4})_")


# --------------------------------------------------------------------------- #
# Reporting
# --------------------------------------------------------------------------- #

class Report:
    """Console reporter that counts warnings (used by --strict)."""

    def __init__(self) -> None:
        self.warnings = 0

    def warn(self, message: str) -> None:
        self.warnings += 1
        print(f"WARNING {message}")

    @staticmethod
    def note(message: str) -> None:
        print(f"NOTE    {message}")


def col_letter(index: int) -> str:
    """0-based column index -> Excel letter (0 -> 'A')."""
    return get_column_letter(index + 1)


# --------------------------------------------------------------------------- #
# Value helpers
# --------------------------------------------------------------------------- #

def decode_ooxml(text: str) -> str:
    return OOXML_ESCAPE_RE.sub(lambda m: chr(int(m.group(1), 16)), text)


def fmt_number(value: float) -> str:
    """1.0 -> '1', 0.5 -> '0.5', 0.25 -> '0.25'."""
    if float(value).is_integer():
        return str(int(value))
    return repr(float(value))


def clean_text(value) -> str:
    """Free text: keep verbatim (only decode OOXML escapes)."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (int, float)):
        return fmt_number(value)
    return decode_ooxml(str(value))


def clean_score(value) -> tuple[str, bool]:
    """
    Evaluation cell -> (csv value, is_numeric).
    Handles real numbers, numbers stored as text (e.g. ' 1\\n', '0,5'),
    and genuinely textual answers (returned trimmed, as-is).
    """
    if value is None:
        return "", True
    if isinstance(value, bool):
        return str(value), False
    if isinstance(value, (int, float)):
        return fmt_number(value), True
    text = decode_ooxml(str(value)).strip()
    if text == "":
        return "", True
    try:
        return fmt_number(float(text.replace(",", "."))), True
    except ValueError:
        return text, False


def parse_question_number(value) -> int | None:
    if value is None:
        return None
    try:
        number = float(str(value).strip().rstrip("."))
    except ValueError:
        return None
    return int(number) if number.is_integer() else None


# --------------------------------------------------------------------------- #
# Header parsing
# --------------------------------------------------------------------------- #

def fold(text: str) -> str:
    """Lower-case and strip accents: 'Justificación' -> 'justificacion'."""
    decomposed = unicodedata.normalize("NFKD", text)
    return "".join(c for c in decomposed if not unicodedata.combining(c)).lower()


def parse_eval_header(header: str) -> str | None:
    """'BO_Evaluation' -> 'BO'; any other header -> None."""
    m = EVAL_HEADER_RE.match(header)
    if m and m.group(2).strip().lower() in EVALUATION_SUFFIXES:
        return m.group(1).upper()
    return None


def find_column(header: list[str], out_name: str, sheet_name: str) -> int:
    """Locate one fixed column by alias; stop if missing or ambiguous."""
    norm = [h.strip().lower() for h in header]
    aliases = FIXED_COLUMNS[out_name]
    matches = [i for i, h in enumerate(norm) if h in aliases]
    if not matches:
        sys.exit(f"ERROR: column '{out_name}' not found in sheet "
                 f"'{sheet_name}'. Header: {header}")
    if len(matches) > 1:
        found = ", ".join(f"{col_letter(i)} {header[i]!r}" for i in matches)
        sys.exit(f"ERROR: column '{out_name}' is ambiguous in sheet "
                 f"'{sheet_name}' - it matches several columns: {found}")
    return matches[0]


def find_fixed_columns(header: list[str], sheet_name: str) -> dict[str, int]:
    """Locate the fixed columns by alias; stop if missing or ambiguous."""
    return {out_name: find_column(header, out_name, sheet_name)
            for out_name in FIXED_COLUMNS}


def find_eval_columns(header: list[str], tag: str, sheet_name: str,
                      report: Report) -> dict[str, int]:
    """
    Return {country_code: column index} of the "XX_Evaluation" columns.
    Stops on duplicated or missing countries; warns about unlisted ones.
    """
    evals: dict[str, list[int]] = {}
    for i, h in enumerate(header):
        code = parse_eval_header(h)
        if code is not None:
            evals.setdefault(code, []).append(i)

    # The same country appearing twice is ambiguous.
    duplicates = {code: idx for code, idx in evals.items() if len(idx) > 1}
    if duplicates:
        details = "; ".join(
            f"{code}_Evaluation in columns "
            f"{', '.join(col_letter(i) for i in idx)}"
            for code, idx in sorted(duplicates.items()))
        sys.exit(f"ERROR: duplicate country column(s) in sheet "
                 f"'{sheet_name}': {details}")

    # Countries in the workbook that are not configured are not exported.
    known = {code for code, _ in COUNTRIES}
    unlisted = sorted(set(evals) - known)
    if unlisted:
        report.warn(f"[{tag}] sheet '{sheet_name}' has evaluation column(s) "
                    f"for country code(s) {unlisted} that are not in "
                    f"COUNTRIES; they are NOT exported.")

    missing = [code for code, _ in COUNTRIES if code not in evals]
    if missing:
        sys.exit(f"ERROR: evaluation column(s) for {missing} not found in "
                 f"sheet '{sheet_name}'. Header: {header}")

    return {code: evals[code][0] for code, _ in COUNTRIES}


def find_country_columns(header: list[str], lang: str, sheet_name: str,
                         fixed_idx: dict[str, int],
                         report: Report) -> dict[str, dict[str, int]]:
    """
    Return {country_code: {"Evaluation": i, "Justification": i, "Sources": i}}.

    Evaluation columns are found by their "XX_Evaluation" header. The two
    columns right after each one are the Justification and the Sources,
    identified by keyword (in either order).
    """
    evals = find_eval_columns(header, lang, sheet_name, report)

    columns: dict[str, dict[str, int]] = {}
    for code, _ in COUNTRIES:
        eval_i = evals[code]
        # Country blocks: Evaluation, then Justification and Sources in the
        # next two columns (either order, identified by keyword).
        block = [eval_i + k for k in range(1, len(BLOCK_KEYWORDS) + 1)]
        if block[-1] >= len(header):
            sys.exit(f"ERROR: sheet '{sheet_name}': {code}_Evaluation is in "
                     f"column {col_letter(eval_i)}, but there aren't "
                     f"{len(block)} columns after it for the Justification "
                     f"and Sources.")
        seen = ", ".join(f"{col_letter(p)} {header[p]!r}" for p in block)

        cols = {"Evaluation": eval_i}
        for kind, keywords in BLOCK_KEYWORDS.items():
            hits = [p for p in block
                    if any(k in fold(header[p]) for k in keywords)]
            if len(hits) != 1:
                problem = "none" if not hits else "more than one"
                sys.exit(f"ERROR: sheet '{sheet_name}': can't identify the "
                         f"{code} {kind} column - {problem} of the columns "
                         f"after {code}_Evaluation ({col_letter(eval_i)}) "
                         f"contains {list(keywords)}. Found: {seen}")
            cols[kind] = hits[0]
        columns[code] = cols

    # Final safety net: no column may feed two different output fields.
    owners: dict[int, list[str]] = {}
    for name, i in fixed_idx.items():
        owners.setdefault(i, []).append(name)
    for code, cols in columns.items():
        for kind, i in cols.items():
            owners.setdefault(i, []).append(f"{code}_{kind}")
    clashes = {i: names for i, names in owners.items() if len(names) > 1}
    if clashes:
        details = "; ".join(f"{col_letter(i)} -> {', '.join(names)}"
                            for i, names in sorted(clashes.items()))
        sys.exit(f"ERROR: column(s) used for more than one field in sheet "
                 f"'{sheet_name}': {details}")

    return columns


# --------------------------------------------------------------------------- #
# Sheet reading
# --------------------------------------------------------------------------- #

def read_scores(wb, sheet_name: str,
                report: Report) -> dict[int, dict[str, str]]:
    """
    Return {question_number: {"<Country>_Evaluation": score}} from the
    scores sheet (the source of truth for every CSV).
    """
    if sheet_name not in wb.sheetnames:
        sys.exit(f"ERROR: scores sheet '{sheet_name}' not found. "
                 f"Available sheets: {wb.sheetnames}")

    rows = wb[sheet_name].iter_rows(values_only=True)
    header = [("" if h is None else str(h)) for h in next(rows)]

    # Only the question number and the evaluation columns are needed here;
    # the texts come from the language sheets.
    q_idx = find_column(header, "Question Number", sheet_name)
    eval_idx = find_eval_columns(header, sheet_name, sheet_name, report)

    allowed = ", ".join(sorted(ALLOWED_SCORES, key=float))
    scores: dict[int, dict[str, str]] = {}
    for row_number, row in enumerate(rows, start=2):
        row = list(row) + [None] * (len(header) - len(row))
        q = parse_question_number(row[q_idx])
        if q is None:
            continue  # blank / footer rows
        if q in scores:
            report.warn(f"[{sheet_name}] duplicate question {q} at row "
                        f"{row_number}; keeping the first one.")
            continue

        record: dict[str, str] = {}
        for code, country in COUNTRIES:
            score, numeric = clean_score(row[eval_idx[code]])
            if not numeric:
                report.note(f"[{sheet_name}] Q{q} {country}: textual "
                            f"evaluation {score!r}")
            elif score and score not in ALLOWED_SCORES:
                report.warn(f"[{sheet_name}] Q{q} {country}: unexpected score "
                            f"{score!r} (row {row_number}; allowed: {allowed})")
            record[f"{country}_Evaluation"] = score
        scores[q] = record

    if not scores:
        sys.exit(f"ERROR: no question rows found in sheet '{sheet_name}'.")
    return scores


def read_sheet(wb, lang: str, sheet_name: str,
               report: Report) -> dict[int, dict[str, str]]:
    """
    Return {question_number: {csv_column: value}} for one language sheet.

    The "<Country>_Evaluation" values read here are the sheet's own scores,
    used only for the cross-check; the exported scores come from SCORES_SHEET.
    """
    if sheet_name not in wb.sheetnames:
        sys.exit(f"ERROR: sheet '{sheet_name}' ({lang}) not found. "
                 f"Available sheets: {wb.sheetnames}")

    rows = wb[sheet_name].iter_rows(values_only=True)
    header = [("" if h is None else str(h)) for h in next(rows)]

    # Fixed columns
    fixed_idx = find_fixed_columns(header, sheet_name)

    # Country blocks: Evaluation (by header), Justification and Sources (by
    # keyword, in the two columns after it)
    country_cols = find_country_columns(header, lang, sheet_name, fixed_idx, report)

    data: dict[int, dict[str, str]] = {}
    for row_number, row in enumerate(rows, start=2):
        row = list(row) + [None] * (len(header) - len(row))
        q = parse_question_number(row[fixed_idx["Question Number"]])
        if q is None:
            continue  # blank / footer rows
        if q in data:
            report.warn(f"[{lang}] duplicate question {q} at row "
                        f"{row_number}; keeping the first one.")
            continue

        record: dict[str, str] = {
            "Dimension": clean_text(row[fixed_idx["Dimension"]]),
            "Categories": clean_text(row[fixed_idx["Categories"]]),
            "Category ID": clean_text(row[fixed_idx["Category ID"]]).strip(),
            "Question Number": str(q),
            "Question": clean_text(row[fixed_idx["Question"]]),
        }
        for code, country in COUNTRIES:
            cols = country_cols[code]
            # Sheet's own score: only compared against SCORES_SHEET.
            score, _ = clean_score(row[cols["Evaluation"]])
            record[f"{country}_Evaluation"] = score
            record[f"{country}_Justification"] = clean_text(row[cols["Justification"]])
            record[f"{country}_Sources"] = clean_text(row[cols["Sources"]])
        data[q] = record

    if not data:
        sys.exit(f"ERROR: no question rows found in sheet '{sheet_name}'.")
    return data


# --------------------------------------------------------------------------- #
# Checks
# --------------------------------------------------------------------------- #

def cross_check(all_data: dict[str, dict[int, dict[str, str]]],
                scores: dict[int, dict[str, str]],
                report: Report) -> None:
    """
    Warn about questions / scores in the language sheets that differ from
    SCORES_SHEET, and about Category IDs that differ from REFERENCE_LANG.
    """
    ref = all_data[REFERENCE_LANG]
    for lang, data in all_data.items():
        only_scores = sorted(set(scores) - set(data))
        only_lang = sorted(set(data) - set(scores))
        if only_scores:
            report.warn(f"[{lang}] questions missing (present in "
                        f"{SCORES_SHEET}): {only_scores}")
        if only_lang:
            report.warn(f"[{lang}] extra questions (absent in "
                        f"{SCORES_SHEET}, exported with blank scores): "
                        f"{only_lang}")

        if lang != REFERENCE_LANG:
            for q in sorted(set(ref) & set(data)):
                if ref[q]["Category ID"] != data[q]["Category ID"]:
                    report.warn(f"[{lang}] Q{q} Category ID "
                                f"{data[q]['Category ID']!r} != {REFERENCE_LANG} "
                                f"{ref[q]['Category ID']!r}")

        for q in sorted(set(scores) & set(data)):
            for _, country in COUNTRIES:
                col = f"{country}_Evaluation"
                if scores[q][col] != data[q][col]:
                    report.warn(f"[{lang}] Q{q} {country}: score "
                                f"{data[q][col]!r} but {SCORES_SHEET} has "
                                f"{scores[q][col]!r}")


# --------------------------------------------------------------------------- #
# Output
# --------------------------------------------------------------------------- #

def csv_columns() -> list[str]:
    cols = ["Dimension", "Categories", "Category ID", "Question Number", "Question"]
    for _, country in COUNTRIES:
        cols += [f"{country}_Evaluation",
                 f"{country}_Justification",
                 f"{country}_Sources"]
    return cols


def write_csv(path: Path, data: dict[int, dict[str, str]],
              reference: dict[int, dict[str, str]],
              scores: dict[int, dict[str, str]]) -> None:
    columns = csv_columns()
    with path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.writer(fh, lineterminator="\n", quoting=csv.QUOTE_MINIMAL)
        writer.writerow(columns)
        for q in sorted(data):
            record = dict(data[q])
            # Dimension / Categories always in English
            if q in reference:
                record["Dimension"] = reference[q]["Dimension"]
                record["Categories"] = reference[q]["Categories"]
            # Scores always from SCORES_SHEET (blank if the question is not
            # there; already reported by cross_check)
            for _, country in COUNTRIES:
                col = f"{country}_Evaluation"
                record[col] = scores.get(q, {}).get(col, "")
            writer.writerow([record[c] for c in columns])


# --------------------------------------------------------------------------- #
# Main
# --------------------------------------------------------------------------- #

def main() -> None:
    parser = argparse.ArgumentParser(
        description="Split the mining scorecard workbook into EN/ES/PT CSVs.")
    parser.add_argument("xlsx", type=Path, help="Path to the .xlsx workbook")
    parser.add_argument("-o", "--outdir", type=Path, default=DEFAULT_OUTDIR,
                        help="Output directory (default: "
                             "public/amazon-mining-policy-scoreboard/data/ "
                             "in the repo root)")
    parser.add_argument("--prefix", default="policy-scoreboard-data",
                        help="Output file prefix (default: policy-scoreboard-data)")
    parser.add_argument("--strict", action="store_true",
                        help="Exit with an error code if any warning is found "
                             "(while reading or cross-checking)")
    args = parser.parse_args()

    if not args.xlsx.exists():
        sys.exit(f"ERROR: file not found: {args.xlsx}")

    report = Report()

    # read_only=True: the workbook contains a pivot cache that openpyxl cannot
    # parse in normal mode. data_only=True: use cached values, not formulas.
    wb = openpyxl.load_workbook(args.xlsx, read_only=True, data_only=True)
    try:
        scores = read_scores(wb, SCORES_SHEET, report)
        all_data = {lang: read_sheet(wb, lang, name, report)
                    for lang, name in SHEETS.items()}
    finally:
        wb.close()

    cross_check(all_data, scores, report)

    args.outdir.mkdir(parents=True, exist_ok=True)
    reference = all_data[REFERENCE_LANG]
    for lang, data in all_data.items():
        out = args.outdir / f"{args.prefix}_{lang}.csv"
        write_csv(out, data, reference, scores)
        print(f"Wrote {out}  ({len(data)} questions)")

    if report.warnings:
        print(f"\n{report.warnings} warning(s) - see above.")
        if args.strict:
            sys.exit(1)


if __name__ == "__main__":
    main()